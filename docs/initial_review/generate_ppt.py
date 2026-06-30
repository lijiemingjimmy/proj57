# -*- coding: utf-8 -*-
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("proj57_initial_review_slides.pptx")
BOARD_IMG = Path(r"C:\Users\Laptop\AppData\Local\Temp\codex-clipboard-D8spvi.png")

FONT = "Microsoft YaHei"
NAVY = RGBColor(25, 38, 64)
BLUE = RGBColor(35, 111, 183)
GREEN = RGBColor(35, 132, 87)
ORANGE = RGBColor(214, 111, 43)
GRAY = RGBColor(92, 99, 112)
LIGHT = RGBColor(245, 247, 250)
WHITE = RGBColor(255, 255, 255)
LINE = RGBColor(219, 225, 232)


def set_font(run, size=18, bold=False, color=NAVY):
    run.font.name = FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_textbox(slide, x, y, w, h, text="", size=18, color=NAVY, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    if align is not None:
        p.alignment = align
    r = p.add_run()
    r.text = text
    set_font(r, size=size, bold=bold, color=color)
    return box


def add_header(slide, title, page, prs):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), prs.slide_width, Inches(0.18))
    bar.fill.solid()
    bar.fill.fore_color.rgb = BLUE
    bar.line.fill.background()
    add_textbox(slide, 0.55, 0.38, 10.8, 0.55, title, 25, NAVY, True)
    add_textbox(slide, 11.85, 0.45, 0.85, 0.35, f"{page}/12", 10, GRAY, False, PP_ALIGN.RIGHT)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.55), Inches(0.98), Inches(12.2), Inches(0.012))
    line.fill.solid()
    line.fill.fore_color.rgb = LINE
    line.line.fill.background()


def add_bullets(slide, items, x=0.72, y=1.32, w=11.9, h=5.55, size=19, color=NAVY):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.color.rgb = color
        p.space_after = Pt(8)
    return box


def add_tag(slide, x, y, text, color=GREEN):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.55), Inches(0.46))
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = Inches(0.12)
    tf.margin_right = Inches(0.12)
    p = tf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run()
    r.text = text
    set_font(r, 13, True, WHITE)
    return shape


def add_panel(slide, x, y, w, h, title, body, accent=BLUE):
    panel = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    panel.fill.solid()
    panel.fill.fore_color.rgb = LIGHT
    panel.line.color.rgb = LINE
    stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
    stripe.fill.solid()
    stripe.fill.fore_color.rgb = accent
    stripe.line.fill.background()
    add_textbox(slide, x + 0.22, y + 0.16, w - 0.35, 0.35, title, 15, NAVY, True)
    add_textbox(slide, x + 0.22, y + 0.58, w - 0.35, h - 0.75, body, 13, GRAY, False)


def add_picture_fit(slide, img_path, x, y, w, h):
    if not img_path.exists():
        add_panel(slide, x, y, w, h, "硬件照片", "板子照片文件未找到，录屏时展示实物即可。", ORANGE)
        return
    from PIL import Image

    with Image.open(img_path) as im:
        iw, ih = im.size
    box_ratio = w / h
    img_ratio = iw / ih
    if img_ratio > box_ratio:
        pic_w = w
        pic_h = w / img_ratio
    else:
        pic_h = h
        pic_w = h * img_ratio
    px = x + (w - pic_w) / 2
    py = y + (h - pic_h) / 2
    slide.shapes.add_picture(str(img_path), Inches(px), Inches(py), width=Inches(pic_w), height=Inches(pic_h))


def build():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    s = prs.slides.add_slide(blank)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = NAVY
    add_textbox(s, 0.8, 0.85, 11.7, 0.9, "Proj57 初赛演示材料", 36, WHITE, True)
    add_textbox(s, 0.82, 1.85, 11.4, 0.65, "面向具身智能的国产操作系统 ACT 模型适配与实时推理优化", 22, WHITE)
    add_tag(s, 0.86, 2.7, "Task3 已有保底", BLUE)
    add_tag(s, 3.62, 2.7, "Task2 RK3588 已跑通", GREEN)
    add_tag(s, 6.38, 2.7, "下一步工程化", ORANGE)
    add_textbox(s, 0.88, 4.92, 11.6, 0.45, "演示主线：阶段总览、真实证据、性能数据、当前风险、下一步计划", 18, WHITE)
    add_textbox(s, 0.88, 5.5, 11.3, 0.4, "核心结论：ONNX -> RKNN -> Orange Pi RK3588 C++ runtime 已形成最小闭环", 17, WHITE)

    s = prs.slides.add_slide(blank)
    add_header(s, "赛题任务与当前策略", 2, prs)
    add_panel(s, 0.7, 1.35, 3.75, 1.85, "任务三", "QEMU / StarryOS CPU 跑通 ACT 推理，是初赛保底链路。", BLUE)
    add_panel(s, 4.78, 1.35, 3.75, 1.85, "任务二", "Orange Pi / RK3588 / NPU 跑通 ACT 推理，是当前硬件适配主线。", GREEN)
    add_panel(s, 8.86, 1.35, 3.75, 1.85, "任务一", "SG2002 / 256 MB 资源约束最强，作为后续高难优化方向。", ORANGE)
    add_bullets(
        s,
        [
            "当前策略：任务三结果保底，任务二 RKNN/NPU 已完成最小闭环。",
            "初赛展示重点：不是只展示想法，而是展示真实日志、模型文件、板端输出和性能数据。",
            "后续路线：runtime 版本对齐、端到端 JPEG 输入、更多样例回归、资源优化。",
        ],
        0.95,
        3.72,
        11.25,
        2.3,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "阶段总览", 3, prs)
    add_bullets(
        s,
        [
            "已完成：模型权重下载、ONNX 导出、ONNX Runtime 正确性验证、本地 benchmark。",
            "已完成：Orange Pi 5 Plus 板端登录、系统信息采集、NPU 节点与 RKNN runtime 探测。",
            "已完成：ONNX 到 RKNN 转换，生成 build/act_rk3588_fp.rknn。",
            "已完成：C++ RKNN runtime 上板运行，两个代表样例方向正确。",
            "正在补齐：视频脚本、性能表、技术文档和可复现材料同步。",
        ],
        0.9,
        1.35,
        11.4,
        4.9,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "模型与推理链路", 4, prs)
    add_panel(s, 0.72, 1.25, 5.85, 1.6, "输入", "image [1,3,224,224]\nstate [1,2]\nlatent [1,32]", BLUE)
    add_panel(s, 6.78, 1.25, 5.85, 1.6, "输出", "action [1,8,3]\n每步包含 left_vel / right_vel / gripper_target", GREEN)
    add_bullets(
        s,
        [
            "ONNX 链路：model.pt -> act.onnx -> ONNX Runtime -> action -> LEFT/RIGHT。",
            "RKNN 链路：act.onnx -> act_rk3588_fp.rknn -> librknnrt -> RK3588/NPU -> action -> LEFT/RIGHT。",
            "方向判断：left_vel < right_vel 为 LEFT，left_vel > right_vel 为 RIGHT。",
        ],
        0.9,
        3.35,
        11.4,
        2.6,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "真实证据 1：ONNX 正确性", 5, prs)
    add_bullets(
        s,
        [
            "已生成 ONNX 模型：build/act.onnx，202,633,726 bytes，约 193.25 MB。",
            "校验命令：python tools/check_onnx.py。",
            "frame_000000.jpg -> LEFT，符合预期。",
            "frame_000227.jpg -> RIGHT，符合预期。",
            "证据文件：artifacts/onnx/check_output.txt、docs/initial_review/baseline_result.md。",
        ],
        0.9,
        1.35,
        11.4,
        4.8,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "真实证据 2：RK3588 板端环境", 6, prs)
    add_bullets(
        s,
        [
            "开发板：Orange Pi 5 Plus，RK3588，aarch64。",
            "系统：Ubuntu 22.04.5 LTS，Rockchip 5.10 内核。",
            "内存：约 3.8 GiB，可支撑任务二验证。",
            "已发现 NPU DRM 节点和 RKNPU 内核初始化日志。",
            "已确认 /usr/lib/librknnrt.so 与 RKNN 头文件存在。",
        ],
        0.75,
        1.3,
        6.4,
        4.9,
        19,
    )
    add_picture_fit(s, BOARD_IMG, 7.55, 1.2, 4.75, 5.35)

    s = prs.slides.add_slide(blank)
    add_header(s, "真实证据 3：RKNN/NPU 运行结果", 7, prs)
    add_bullets(
        s,
        [
            "RKNN 模型：build/act_rk3588_fp.rknn，101,709,294 bytes，约 97.0 MiB。",
            "C++ runtime：tools/rknn_runtime_infer.cpp。",
            "frame_000000：LEFT，rknn_perf_run_us=34968，wall_ms=34.998。",
            "frame_000227：RIGHT，rknn_perf_run_us=25340，wall_ms=25.364。",
            "证据文件：artifacts/rknn/runtime_log_frame000_nhwc.txt、runtime_log_frame227_nhwc.txt。",
        ],
        0.9,
        1.35,
        11.4,
        4.8,
        20,
    )
    add_tag(s, 0.96, 6.23, "任务二最小闭环已跑通", GREEN)

    s = prs.slides.add_slide(blank)
    add_header(s, "性能与资源数据表", 8, prs)
    table = s.shapes.add_table(4, 6, Inches(0.55), Inches(1.35), Inches(12.25), Inches(3.0)).table
    headers = ["平台", "后端", "状态", "模型大小", "推理耗时", "结果"]
    data = [
        ["StarryOS/QEMU", "ONNX CPU", "已有报告", "193.2 MB", "18.4-18.7 s", "LEFT/RIGHT 正确"],
        ["Windows 本机", "ONNX CPU", "已复现", "193.25 MB", "25.91-26.78 ms", "LEFT/RIGHT 正确"],
        ["Orange Pi RK3588", "RKNN/NPU", "已跑通", "97.0 MiB", "25.34-35.00 ms", "LEFT/RIGHT 正确"],
    ]
    for i, width in enumerate([2.15, 1.7, 1.55, 1.85, 2.05, 2.95]):
        table.columns[i].width = Inches(width)
    for j, header in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = BLUE
        for p in cell.text_frame.paragraphs:
            p.font.name = FONT
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = WHITE
    for r, row in enumerate(data, 1):
        for c, value in enumerate(row):
            cell = table.cell(r, c)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if r % 2 else LIGHT
            for p in cell.text_frame.paragraphs:
                p.font.name = FONT
                p.font.size = Pt(11)
                p.font.color.rgb = NAVY
    add_bullets(
        s,
        [
            "RKNN 结果当前只覆盖两个代表样例，后续需要扩展回归集和稳定性统计。",
            "板端 runtime 有版本警告，因此视频中应展示关键行并主动解释。",
        ],
        0.8,
        4.75,
        11.7,
        1.1,
        17,
        GRAY,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "当前工程产物", 9, prs)
    add_bullets(
        s,
        [
            "模型权重：output/train/model.pt，202,962,639 bytes。",
            "ONNX 模型：build/act.onnx，202,633,726 bytes。",
            "RKNN 模型：build/act_rk3588_fp.rknn，101,709,294 bytes。",
            "RKNN 转换脚本：tools/convert_act_rknn.py。",
            "RKNN runtime：tools/rknn_runtime_infer.cpp。",
            "初审材料：技术文档、进度报告、演示脚本、性能表、PPT、PDF。",
        ],
        0.9,
        1.35,
        11.4,
        4.8,
        19,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "当前问题与风险", 10, prs)
    add_bullets(
        s,
        [
            "板端 RKNN runtime API 1.5.2 与 RKNN Toolkit2 2.3.2 不完全匹配。",
            "运行时会打印 invalid run task counter 警告，但仍能返回输出和耗时。",
            "当前 C++ runtime 使用预处理后的 tensor bin，尚未直接支持 JPEG 输入。",
            "当前只验证两个代表样例，仍需更多样例和长时间稳定性测试。",
            "后续任务一 SG2002 内存只有 256 MB，需要模型裁剪、量化或更激进的算子拆分。",
        ],
        0.9,
        1.35,
        11.4,
        4.8,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "初赛演示视频安排", 11, prs)
    add_panel(s, 0.72, 1.25, 2.25, 1.25, "20s", "开场：赛题目标、三档任务、当前路线。", BLUE)
    add_panel(s, 3.16, 1.25, 2.25, 1.25, "60s", "硬件：Orange Pi、SSH、NPU 节点。", GREEN)
    add_panel(s, 5.60, 1.25, 2.25, 1.25, "70s", "基线：ONNX 正确性与 Task3 输出。", BLUE)
    add_panel(s, 8.04, 1.25, 2.25, 1.25, "120s", "任务二：RKNN 输出、耗时、日志证据。", GREEN)
    add_panel(s, 10.48, 1.25, 2.25, 1.25, "40s", "结尾：风险边界和下一步计划。", ORANGE)
    add_bullets(
        s,
        [
            "录制时只展示 RKNN 日志关键行，避免警告刷屏影响观感。",
            "必须主动说明：版本警告存在，但方向、输出和耗时均已返回。",
            "不要把 Task3 说成现场完整跑 QEMU；更稳妥的说法是展示已有报告输出。",
        ],
        0.9,
        3.35,
        11.4,
        2.4,
        20,
    )

    s = prs.slides.add_slide(blank)
    add_header(s, "总结", 12, prs)
    add_bullets(
        s,
        [
            "任务三链路已有正确性和性能证据，可作为初赛保底。",
            "任务二已完成 RK3588/RKNN/NPU 最小闭环。",
            "当前重点从“能不能跑”转为“runtime 版本对齐、稳定性和工程化”。",
            "初赛展示重点：真实证据清晰、数据表完整、问题边界明确、后续路径可执行。",
        ],
        0.9,
        1.45,
        11.4,
        4.4,
        22,
    )
    add_tag(s, 0.96, 6.1, "ONNX 正确性已复现", BLUE)
    add_tag(s, 3.72, 6.1, "RKNN 板端已跑通", GREEN)
    add_tag(s, 6.48, 6.1, "下一步工程化优化", ORANGE)

    prs.save(OUT)
    print(f"generated {OUT} slides={len(prs.slides)} size={OUT.stat().st_size}")


if __name__ == "__main__":
    build()
